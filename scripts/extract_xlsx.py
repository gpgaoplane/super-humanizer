



























































import io
import posixpath
import re
import warnings
import zipfile
from pathlib import Path

from scripts import office
from scripts.records import Document, Unit

FORMAT = "xlsx"
PACKAGE = "openpyxl"


FLOOR = "3.1"



LXML_PACKAGE = "lxml"
LXML_FLOOR = "4"


CANNOT_SEE = (
    "formula-built cell text when no cached value is stored",
    "a chart's series names and category labels, which reference cells the walk already "
    "reads as cells",
    "chart text outside the chart XML openpyxl reads, such as data labels and drawing text "
    "boxes",
    "a data validation's prompt and error text",
    "a comment's author",
    "a threaded comment as Excel shows it, since openpyxl reads the legacy comment written "
    "beside it, and that comment's placeholder paragraph is read as the comment's prose",
    "an image and any text inside it",
    "a pivot table's captions",
    "the side-by-side copy's own gap, since openpyxl writes formulas without their cached "
    "values, so a run on the copy sees none of the formula-built text and Excel recalculates "
    "on open",
    "the side-by-side copy's own losses, since openpyxl's save drops images when Pillow is "
    "absent, drawing shapes and text boxes, data validation and conditional formatting "
    "extensions, sparklines and slicers, and resets the size and place of each comment box",
)


COPY_LINES = (CANNOT_SEE[-2], CANNOT_SEE[-1])



_HEADERS = (("oddHeader", "odd header"), ("oddFooter", "odd footer"),
            ("evenHeader", "even header"), ("evenFooter", "even footer"),
            ("firstHeader", "first header"), ("firstFooter", "first footer"))
_PARTS = ("left", "center", "right")


def is_machine_key(text):

    return text.count("|") >= 2 and " | " not in text


def _rich(title):



    if title is None or title.tx is None or title.tx.rich is None:
        return ""
    return "\n".join("".join(run.t or "" for run in paragraph.r or [])
                     for paragraph in title.tx.rich.p or [])


def _chart_titles(chart):



    axes = []
    for plot in getattr(chart, "_charts", None) or [chart]:
        for name in ("x_axis", "y_axis", "z_axis"):
            axis = getattr(plot, name, None)
            if axis is not None and not any(axis is seen for seen in axes):
                axes.append(axis)
    titled = [text for text in (_rich(axis.title) for axis in axes) if text.strip()]
    title = _rich(chart.title)
    entries = [("chart", None, title)] if title.strip() else []
    return entries + [("axis", index, text) for index, text in enumerate(titled)]


class _Walk:


    def __init__(self):
        self.units, self.protected = [], []
        self.formulas = self.uncached = self.errors = self.equals = 0

    def _add(self, kind, surface, text, locator, frame, weighs=True):
        unit = Unit(id=f"u{len(self.units) + 1:04d}", kind=kind, surface=surface, text=text,
                    locator=locator, weighs=weighs, frame=frame)
        self.units.append(unit)
        return unit

    def sheet(self, sheet, values):


        name = sheet.title
        hidden = sheet.sheet_state != "visible"
        shown = f"{name} (hidden)" if hidden else name
        base = {"sheet": name, "hidden": True} if hidden else {"sheet": name}
        self._add("label", f"sheet {shown} name", name, {**base, "name": True},
                  {"kind": "sheet_name", "writable": False})
        commented = []
        for row in sheet.iter_rows():
            for cell in row:
                self._cell(cell, values, shown, base)
                if cell.comment is not None:
                    commented.append(cell)
        for cell in commented:
            text = cell.comment.text or ""
            if text.strip():
                self._add("notes", f"sheet {shown} comment on {cell.coordinate}", text,
                          {**base, "cell": cell.coordinate, "comment": True},
                          {"kind": "cell_comment", "writable": True})
        for which, words in _HEADERS:
            item = getattr(sheet, which)
            for part in _PARTS:
                text = getattr(item, part).text or ""
                if text.strip():
                    self._add("label", f"sheet {shown} {words} {part}", text,
                              {**base, "header": which, "part": part},
                              {"kind": "header_footer", "writable": True})
        for number, chart in enumerate(sheet._charts):
            for role, index, text in _chart_titles(chart):
                locator = {**base, "chart": number, "title": role}
                surface = f"sheet {shown} chart {number + 1} title"
                if index is not None:
                    locator["index"] = index
                    surface = f"sheet {shown} chart {number + 1} axis title {index + 1}"
                self._add("label", surface, text, locator,
                          {"kind": "chart_text", "writable": False})

    def _cell(self, cell, values, shown, base):

        if cell.data_type == "f":
            self._formula(cell, values, shown, base)
            return
        if cell.data_type == "e":
            return
        value = cell.value
        if isinstance(value, str) and value.startswith("="):
            self.equals += 1
            return
        if not isinstance(value, str) or not value.strip():
            return
        self._string(f"sheet {shown} cell {cell.coordinate}", value.strip(),
                     {**base, "cell": cell.coordinate},
                     {"kind": "cell", "formula": False, "writable": True})

    def _string(self, surface, text, locator, frame):


        key = is_machine_key(text)
        unit = self._add("table_cell", surface, text, locator, frame, weighs=not key)
        if key:
            self.protected.append((unit.id, 0, len(text), "C143"))

    def _formula(self, cell, values, shown, base):


        self.formulas += 1
        formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", "")
        self.protected.append((None, 0, len(formula or ""), "C143"))
        cached = values.cell(row=cell.row, column=cell.column)
        if cached.value is None:
            if cached.data_type != "str":
                self.uncached += 1
            return
        if cached.data_type == "e":
            self.errors += 1
            return
        if not isinstance(cached.value, str) or not cached.value.strip():
            return
        self._string(f"sheet {shown} cell {cell.coordinate} (formula result)",
                     cached.value.strip(), {**base, "cell": cell.coordinate},
                     {"kind": "cell", "formula": True, "writable": False})


def extract(path):











    path = Path(path)
    try:
        from openpyxl import load_workbook
    except ImportError:
        return Document(path=str(path), format=FORMAT, units=[],
                        cannot_see=[office.missing_package(PACKAGE, FORMAT)], encoding="utf-8",
                        newline="\n")
    try:
        from openpyxl.cell.rich_text import CellRichText
        from openpyxl.workbook.properties import CalcProperties
    except ImportError:
        return Document(path=str(path), format=FORMAT, units=[],
                        cannot_see=[office.below_floor(PACKAGE, FLOOR, FORMAT)],
                        encoding="utf-8", newline="\n")
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            book = load_workbook(path)
            values = load_workbook(path, data_only=True)
    except Exception as error:
        raise office.OfficeError(f"{path}: {PACKAGE} could not open the file, "
                                 f"{type(error).__name__}: {error}") from error
    walk = _Walk()
    lines = [f"openpyxl warned while loading the file: {message}"
             for message in dict.fromkeys(str(warning.message) for warning in caught)]
    for name in book.sheetnames:
        sheet = book[name]
        if not hasattr(sheet, "iter_rows"):
            lines.append(f"the chartsheet {name} holds a chart alone, so its text was not read")
            continue
        walk.sheet(sheet, values[name])
    if walk.uncached:
        lines.append(f"{walk.uncached} of {walk.formulas} formula cells carry no cached value, "
                     "so the text they build is unseen")
    if walk.errors:
        lines.append(f"{walk.errors} of {walk.formulas} formula cells hold an error value as "
                     "their cached result, so they are no unit")
    if walk.equals:
        lines.append(f"{walk.equals} string cells start with an equals sign, which the "
                     "typed-string rule passes over, so their text is unseen")
    return Document(path=str(path), format=FORMAT, units=walk.units,
                    cannot_see=list(CANNOT_SEE) + lines, encoding="utf-8", newline="\n",
                    protected=walk.protected)



CHANGED = ("the text differs from the one the extract read, so the workbook changed between "
           "extract and write")
ABSENT = "the sheet, cell or comment the locator names is absent from the workbook"
FORMULA_START = ("the rewrite starts with an equals sign, which a workbook cell reads as a "
                 "formula")


FLATTENED = ("the rewrite of {surface} was written as plain text, so the rich text runs of its "
             "{kind} were flattened")


_EMPTIED = {"cell": "cell", "cell_comment": "comment", "header_footer": "header or footer part"}


_CODE = re.compile("&[PNDTGAFZLCR&]")





_KEPT_CONTROLS = (chr(9), chr(10), chr(13))
_SURROGATES = (chr(0xD800), chr(0xDFFF))
_NONCHARACTERS = (chr(0xFFFE), chr(0xFFFF))


def _unwritable(char):




    first, last = _SURROGATES
    return ((char < " " and char not in _KEPT_CONTROLS) or first <= char <= last
            or char in _NONCHARACTERS)


def _codes(text):

    return " ".join(_CODE.findall(text)) or "none"


def _check_unit(book, unit, text, rich, errors):









    frame, locator = unit.frame, unit.locator
    if not frame.get("writable"):
        return office.READ_ONLY, None
    if not text.strip():
        return f"the rewrite empties the {_EMPTIED[frame['kind']]}", None
    held = next((char for char in text if _unwritable(char)), None)
    if held is not None:
        kind = "control character" if held < " " else "character"
        return (f"the rewrite holds the {kind} U+{ord(held):04X}, which cannot be written "
                f"into the {_EMPTIED[frame['kind']]}"), None
    if locator["sheet"] not in book.sheetnames:
        return ABSENT, None
    sheet = book[locator["sheet"]]
    if frame["kind"] == "header_footer":
        part = getattr(getattr(sheet, locator["header"]), locator["part"])
        if (part.text or "") != unit.text:
            return CHANGED, None
        if _CODE.findall(text) != _CODE.findall(unit.text):
            return (f"the rewrite changes the part's field codes, {_codes(unit.text)} in the "
                    f"extract and {_codes(text)} in the rewrite"), None
        return None, None
    cell = sheet[locator["cell"]]
    if frame["kind"] == "cell_comment":
        if cell.comment is None:
            return ABSENT, None
        if (cell.comment.text or "") != unit.text:
            return CHANGED, None
        return None, None
    value = cell.value
    old = str(value) if isinstance(value, (str, rich)) else None
    if cell.data_type == "f" or old is None or old.strip() != unit.text:
        return CHANGED, None
    new = text.strip()
    if new.startswith("="):
        return FORMULA_START, None
    if new in errors:
        return (f"the rewrite is the error value {new}, which a workbook cell reads as an "
                "error value"), None
    return None, (old[:len(old) - len(old.lstrip())], old[len(old.rstrip()):])


def _write_unit(book, unit, text, rich, errors, lines):





    reason, padding = _check_unit(book, unit, text, rich, errors)
    if reason is not None:
        return reason
    frame, locator = unit.frame, unit.locator
    sheet = book[locator["sheet"]]
    if frame["kind"] == "header_footer":
        getattr(getattr(sheet, locator["header"]), locator["part"]).text = text
    elif frame["kind"] == "cell_comment":
        sheet[locator["cell"]].comment.text = text
    else:
        cell = sheet[locator["cell"]]
        flattened = isinstance(cell.value, rich)
        cell.value = padding[0] + text.strip() + padding[1]
        if flattened:
            lines.append(FLATTENED.format(surface=unit.surface,
                                          kind=_EMPTIED[unit.frame["kind"]]))
    return None


def write_copy(path, document, texts, file_digest):



















    path = Path(path)
    target = office.copy_path(path)
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import ERROR_CODES
        from openpyxl.cell.rich_text import CellRichText
        from openpyxl.workbook.properties import CalcProperties
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    applied, refused, lines = [], [], []
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = load_workbook(io.BytesIO(data), rich_text=True)
        for unit in document.units:
            if unit.id not in texts:
                continue
            reason = _write_unit(book, unit, texts[unit.id], CellRichText, ERROR_CODES, lines)
            if reason is None:
                applied.append(unit.id)
            else:
                refused.append((unit.id, reason))
        if book.calculation is None:
            book.calculation = CalcProperties()
        book.calculation.fullCalcOnLoad = True
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if texts and not applied:
        return None, [], refused
    try:
        office.save_copy(target, book.save)
    except Exception as error:
        return None, [], [(None, f"the copy could not be saved at {target}: "
                                 f"{type(error).__name__}: {error}")]
    document.cannot_see.extend(lines)
    return target, applied, refused





STORED = ("the cell is stored with the type {type}, and the in-place write rewrites a shared "
          "string or an inline string alone")


SHARED_STRINGS_ABSENT = ("the workbook lists no shared strings part for a cell stored as a "
                         "shared string")



THREADED = ("the comment sits on a sheet whose relationships hold a threaded comment, whose "
            "text Excel shows from its own part, so the legacy text openpyxl reads is never "
            "what a reader sees")



HEADER_MERGE = ("the rewrite would read as a header code or as part of the one before it, so "
                "the header's codes would change")


STALE_RESULTS = ("formulas that read a rewritten cell show their old result until Excel "
                 "recalculates, which the file now asks for on open")

_WORKBOOK = "xl/workbook.xml"
_WORKBOOK_RELS = "xl/_rels/workbook.xml.rels"
_RELATIONSHIP_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _entry(source, target):



    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source), target))


class _Parts:



    def __init__(self, etree, package):
        self.etree, self.package = etree, package
        self.names = set(package.namelist())
        self.parser = etree.XMLParser(resolve_entities=False)
        self.trees, self.changed = {}, set()

    def root(self, name):
        if name not in self.trees:
            self.trees[name] = self.etree.fromstring(self.package.read(name), self.parser)
        return self.trees[name]


def _targets(etree, parts):


    rels = {rel.get("Id"): rel for rel in parts.root(_WORKBOOK_RELS)
            if rel.get("TargetMode") != "External"}
    workbook = parts.root(_WORKBOOK)
    main = etree.QName(workbook).namespace
    sheets = {}
    for sheet in workbook.iter(f"{{{main}}}sheet"):
        rel = rels.get(sheet.get(_RELATIONSHIP_ID))
        if rel is not None:
            sheets[sheet.get("name")] = _entry(_WORKBOOK, rel.get("Target"))
    shared = next((_entry(_WORKBOOK, rel.get("Target")) for rel in rels.values()
                   if (rel.get("Type") or "").endswith("/sharedStrings")), None)
    return sheets, shared


def _string_text(main, string):


    return (string.findtext(f"{{{main}}}t") or "") + "".join(
        run.findtext(f"{{{main}}}t") or "" for run in string.findall(f"{{{main}}}r"))


def _one_t(etree, main, string, value):


    for child in list(string):
        string.remove(child)
    held = etree.SubElement(string, f"{{{main}}}t")
    held.text = value
    if value != value.strip():
        held.set(_XML_SPACE, "preserve")


def _write_cell(etree, parts, sheets, shared, unit, value, lines):












    entry = sheets.get(unit.locator["sheet"])
    if entry not in parts.names:
        return ABSENT
    root = parts.root(entry)
    main = etree.QName(root).namespace
    cell = next((cell for cell in root.iter(f"{{{main}}}c")
                 if cell.get("r") == unit.locator["cell"]), None)
    if cell is None:
        return ABSENT
    kind = cell.get("t", "n")
    if kind not in ("s", "inlineStr"):
        return STORED.format(type=kind)
    if kind == "inlineStr":
        string = cell.find(f"{{{main}}}is")
    else:
        if shared not in parts.names:
            return SHARED_STRINGS_ABSENT
        table = parts.root(shared)
        strings = table.findall(f"{{{main}}}si")
        index = (cell.findtext(f"{{{main}}}v") or "").strip()
        string = strings[int(index)] if index.isdigit() and int(index) < len(strings) else None
    if string is None or _string_text(main, string).strip() != unit.text:
        return CHANGED
    if string.find(f"{{{main}}}r") is not None:
        lines.append(FLATTENED.format(surface=unit.surface,
                                      kind=_EMPTIED[unit.frame["kind"]]))
    if kind == "inlineStr":
        _one_t(etree, main, string, value)
    else:
        added = table.makeelement(f"{{{main}}}si", {})
        table.insert(table.index(strings[-1]) + 1, added)
        _one_t(etree, main, added, value)
        unique = table.get("uniqueCount")
        if unique is not None:
            table.set("uniqueCount", str(int(unique) + 1))
        cell.find(f"{{{main}}}v").text = str(len(strings))
        parts.changed.add(shared)
    parts.changed.add(entry)
    return None


def _sheet_rels(parts, entry):


    name = posixpath.join(posixpath.dirname(entry), "_rels",
                          posixpath.basename(entry) + ".rels")
    if name not in parts.names:
        return []
    return [rel for rel in parts.root(name) if rel.get("TargetMode") != "External"]


def _write_comment(etree, parts, sheets, unit, value, lines):











    entry = sheets.get(unit.locator["sheet"])
    if entry not in parts.names:
        return ABSENT
    rels = _sheet_rels(parts, entry)
    if any((rel.get("Type") or "").endswith("threadedComment") for rel in rels):
        return THREADED
    target = next((_entry(entry, rel.get("Target")) for rel in rels
                   if (rel.get("Type") or "").endswith("/comments")), None)
    if target not in parts.names:
        return ABSENT
    root = parts.root(target)
    main = etree.QName(root).namespace
    comment = next((comment for comment in root.iter(f"{{{main}}}comment")
                    if comment.get("ref") == unit.locator["cell"]), None)
    if comment is None:
        return ABSENT
    text = comment.find(f"{{{main}}}text")
    if text is None or _string_text(main, text) != unit.text:
        return CHANGED
    if text.find(f"{{{main}}}r") is not None:
        lines.append(FLATTENED.format(surface=unit.surface,
                                      kind=_EMPTIED[unit.frame["kind"]]))
    _one_t(etree, main, text, value)
    parts.changed.add(target)
    return None



_MARKERS = (("left", "&L"), ("center", "&C"), ("right", "&R"))


def _write_header(etree, names, parts, sheets, unit, value):














    entry = sheets.get(unit.locator["sheet"])
    if entry not in parts.names:
        return ABSENT
    root = parts.root(entry)
    main = etree.QName(root).namespace
    node = root.find(f"{{{main}}}headerFooter/{{{main}}}{unit.locator['header']}")
    if node is None:
        return ABSENT
    held = node.text or ""
    segments = names._split_string(held)
    segment = segments[unit.locator["part"]] or ""
    if names.FORMAT_REGEX.sub("", segment) != unit.text or segment.count(unit.text) != 1:
        return CHANGED
    new = segment.replace(unit.text, value, 1)
    if names.FORMAT_REGEX.sub("", new) != value:
        return HEADER_MERGE
    old = "".join(marker + segments[part] for part, marker in _MARKERS if segments[part])
    segments[unit.locator["part"]] = new
    node.text = "".join(marker + segments[part] for part, marker in _MARKERS
                        if segments[part]) + held[len(old):]
    parts.changed.add(entry)
    return None


def _flag(etree, parts, elements):






    root = parts.root(_WORKBOOK)
    main = etree.QName(root).namespace
    calc = root.find(f"{{{main}}}calcPr")
    if calc is None:
        before = {f"{{{main}}}{name}" for name in elements[:elements.index("calcPr")]}
        place = max((index for index, child in enumerate(root) if child.tag in before),
                    default=-1)
        calc = root.makeelement(f"{{{main}}}calcPr", {})
        root.insert(place + 1, calc)
    elif calc.get("fullCalcOnLoad") in ("1", "true"):
        return False
    calc.set("fullCalcOnLoad", "1")
    return True


def write_in_place(path, document, texts, file_digest):






















    path = Path(path)
    if not texts:
        return None, [], []
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import ERROR_CODES
        from openpyxl.cell.rich_text import CellRichText
        from openpyxl.packaging.workbook import WorkbookPackage
        from openpyxl.worksheet import header_footer
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    try:
        from lxml import etree
    except ImportError:
        return None, [], [(None, office.missing_package(LXML_PACKAGE, FORMAT))]
    checked, refused = [], []
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = load_workbook(io.BytesIO(data), rich_text=True)
        for unit in document.units:
            if unit.id not in texts:
                continue
            text = texts[unit.id]
            reason, padding = _check_unit(book, unit, text, CellRichText, ERROR_CODES)
            if reason is None:
                checked.append((unit, text if padding is None
                                else padding[0] + text.strip() + padding[1]))
            else:
                refused.append((unit.id, reason))
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if not checked:
        return None, [], refused
    applied, lines, replacements, cells = [], [], {}, False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            order = package.namelist()
            parts = _Parts(etree, package)
            sheets, shared = _targets(etree, parts)
            for unit, value in checked:
                kind = unit.frame["kind"]
                if kind == "cell":
                    reason = _write_cell(etree, parts, sheets, shared, unit, value, lines)
                elif kind == "cell_comment":
                    reason = _write_comment(etree, parts, sheets, unit, value, lines)
                else:
                    reason = _write_header(etree, header_footer, parts, sheets, unit, value)
                if reason is None:
                    applied.append(unit.id)
                    cells = cells or kind == "cell"
                else:
                    refused.append((unit.id, reason))
            if cells and _flag(etree, parts, WorkbookPackage.__elements__):
                parts.changed.add(_WORKBOOK)
            for name in parts.changed:
                replacements[name] = office.with_leading_bytes(
                    package.read(name), etree.tostring(parts.root(name), encoding="UTF-8"))
    except Exception as error:
        return None, [], [(None, f"{LXML_PACKAGE} raised {type(error).__name__}: {error}")]
    position = {unit.id: index for index, unit in enumerate(document.units)}
    refused.sort(key=lambda pair: position[pair[0]])
    if not applied:
        return None, [], refused
    kept = {unit_id: texts[unit_id] for unit_id in applied}

    def census(temp):
        return office.reads_as(document, extract(temp), kept)

    reason = office.write_in_place(path, data, replacements, census)
    if reason is not None:
        return None, [], [(None, reason)]
    replaced = [name for name in order if name in replacements]
    document.cannot_see.extend(lines)
    if cells:
        document.cannot_see.append(STALE_RESULTS)
    document.cannot_see.append(office.IN_PLACE_PARTS.format(
        path=path, count=len(replaced), total=len(order), names=", ".join(replaced)))
    return path, applied, refused
