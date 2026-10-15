










































import io
import xml.etree.ElementTree as ElementTree
import zipfile
from collections import Counter
from pathlib import Path

from scripts import office
from scripts.records import Document, Unit

FORMAT = "pptx"
PACKAGE = "python-pptx"

FLOOR = "1.0"


CANNOT_SEE = (
    "images and any text inside them",
    "alt text on a picture or a shape",
    "the text of a slide layout or master",
    "a chart's data labels and any chart text outside its title, axis titles, series names "
    "and categories",
    "a comment part",
    "the text a binding template's own dump would exclude, since the run has no template "
    "input",
    "an embedded workbook behind a chart",
)

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

_TEMPLATE_TYPES = (
    "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml",
    "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
)



_CHART_SURFACES = {"title": "title", "axis": "axis title {}", "series": "series {}",
                   "category": "category {}"}


def _paragraph(element):



    pieces, runs, breaks, fields = [], [], 0, []
    for child in element:
        if child.tag == _A + "r":
            node = child.find(_A + "t")
            text = (node.text or "") if node is not None else ""
            pieces.append(text)
            runs.append(len(text))
        elif child.tag == _A + "br":
            pieces.append("\n")
            breaks += 1
        elif child.tag == _A + "fld":
            node = child.find(_A + "t")
            fields.append((node.text or "") if node is not None else "")
    return "".join(pieces), runs, breaks, fields


def _rich(title):


    return "\n".join(_paragraph(element)[0]
                     for element in title.findall(f"{_C}tx/{_C}rich/{_A}p"))


def _chart_strings(chart):







    strings = []
    if chart.has_title and chart.chart_title.has_text_frame:
        strings.append(("title", None, _rich(chart.chart_title.element)))
    titles = chart.element.findall(f"{_C}chart/{_C}plotArea/*/{_C}title")
    strings.extend(("axis", index, _rich(title)) for index, title in enumerate(titles))
    strings.extend(("series", index, series.name) for index, series in enumerate(chart.series))
    for plot in chart.plots:
        categories = [str(category) for category in plot.categories]
        if any(category.strip() for category in categories):
            strings.extend(("category", index, text) for index, text in enumerate(categories))
            break
    return [entry for entry in strings if entry[2].strip()]


def _shapes(shapes, path=()):


    for shape in shapes:
        here = path + (shape.shape_id,)
        if shape.element.tag == _P + "grpSp":
            yield from _shapes(shape.shapes, here)
        else:
            yield here, shape


class _Walk:


    def __init__(self, titles, furniture):
        self.units, self.protected, self.lines = [], [], []
        self.titles, self.furniture = titles, furniture
        self.counts = Counter()
        self.on_slides = self.reached = 0

    def _next(self, key):

        self.counts[key] += 1
        return self.counts[key]

    def _label(self, prefix, key):

        count = self._next(key)
        return prefix if count == 1 else f"{prefix} {count}"

    def _body(self, body, kind, locator, name, furniture=False, on_slide=True):











        if body is None:
            return
        for index, element in enumerate(body.findall(_A + "p")):
            if on_slide:
                self.reached += _text_nodes(element)
            text, runs, breaks, fields = _paragraph(element)
            unit = None
            if text.strip():
                unit = Unit(
                    id=f"u{len(self.units) + 1:04d}", kind=kind, surface=name(index),
                    text=text, locator={**locator, "paragraph": index},
                    weighs=not (furniture or fields),
                    frame={"kind": "pptx_paragraph", "runs": runs, "breaks": breaks,
                           "fields": len(fields), "writable": True})
                self.units.append(unit)
                if furniture:
                    self.protected.append((unit.id, 0, len(text), "C142"))
            for field in fields:
                self.protected.append((unit.id if unit else None, 0, len(field), "C142"))

    def slide(self, slide, number):
        self.counts = Counter()
        self.on_slides += _text_nodes(ElementTree.fromstring(slide.part.blob))
        shapes = list(_shapes(slide.shapes))
        texts = [(path, shape) for path, shape in shapes if shape.has_text_frame]
        for path, shape in texts:
            if _placeholder(shape) in self.titles:
                self._body(shape.element.find(_P + "txBody"), "title",
                           {"slide": number, "shape": list(path)},
                           lambda index: self._label(f"slide {number} title", "title"))
        for path, shape in texts:
            if _placeholder(shape) not in self.titles:
                self._body(shape.element.find(_P + "txBody"), "paragraph",
                           {"slide": number, "shape": list(path)},
                           lambda index: f"slide {number} paragraph {self._next('paragraph')}",
                           furniture=_placeholder(shape) in self.furniture)
        for path, shape in shapes:
            if shape.has_table:
                self._table(path, shape, number)
        for path, shape in shapes:
            if shape.has_chart:
                self._chart(path, shape, number)
        for path, shape in shapes:
            relids = next(shape.element.iter(_DGM + "relIds"), None)
            if relids is not None:
                self._diagram(path, relids, slide, number)
        if slide.has_notes_slide:
            placeholder = slide.notes_slide.notes_placeholder
            if placeholder is not None:
                self._body(placeholder.element.find(_P + "txBody"), "notes",
                           {"slide": number, "notes": True},
                           lambda index: f"slide {number} notes {self._next('notes')}",
                           on_slide=False)

    def _table(self, path, shape, number):
        name = self._label("table", "table")
        grid = shape.element.find(".//" + _A + "tbl")
        if grid is None:
            self.lines.append(f"slide {number}: a table frame holds no table, so it was not read")
            return
        for row, line in enumerate(grid.findall(_A + "tr")):
            for column, cell in enumerate(line.findall(_A + "tc")):
                self._body(cell.find(_A + "txBody"), "table_cell",
                           {"slide": number, "shape": list(path), "cell": [row, column]},
                           lambda index: (f"slide {number} {name} cell r{row + 1}c{column + 1} "
                                          f"paragraph {self._next('table_cell')}"))


    def _chart(self, path, shape, number):


        name = self._label("chart", "chart")
        try:
            strings = _chart_strings(shape.chart)
        except Exception as error:
            self.lines.append(f"slide {number}: a chart's text was not read, since the chart "
                              f"API raised {type(error).__name__}")
            return
        for role, index, text in strings:
            locator = {"slide": number, "shape": list(path), "chart": role}
            if index is not None:
                locator["index"] = index
            surface = _CHART_SURFACES[role].format(None if index is None else index + 1)
            self.units.append(Unit(
                id=f"u{len(self.units) + 1:04d}", kind="label",
                surface=f"slide {number} {name} {surface}", text=text, locator=locator,
                frame={"kind": "chart_text", "writable": False}))

    def _diagram(self, path, relids, slide, number):



        try:
            part = slide.part.related_part(relids.get(_R + "dm"))
        except (KeyError, ValueError):
            self.lines.append(f"slide {number}: a SmartArt diagram's data part is absent, so "
                              "its text was not read")
            return
        try:
            model = ElementTree.fromstring(part.blob)
        except ElementTree.ParseError:
            self.lines.append(f"slide {number}: a SmartArt diagram's data part could not be "
                              "parsed, so its text was not read")
            return
        index = 0
        for point in model.iter(_DGM + "pt"):
            text = "\n".join(_paragraph(element)[0]
                             for element in point.findall(f"{_DGM}t/{_A}p"))
            if not text.strip():
                continue
            self.units.append(Unit(
                id=f"u{len(self.units) + 1:04d}", kind="label",
                surface=f"slide {number} diagram text {self._next('diagram')}", text=text,
                locator={"slide": number, "shape": list(path), "diagram": index},
                frame={"kind": "diagram_text", "writable": False}))
            index += 1


def _placeholder(shape):

    return shape.placeholder_format.type if shape.is_placeholder else None


def _text_nodes(element):

    return sum(1 for node in element.iter(_A + "t") if (node.text or "").strip())


def _template_text(deck):


    return sum(_text_nodes(ElementTree.fromstring(part.blob))
               for part in deck.part.package.iter_parts()
               if part.content_type in _TEMPLATE_TYPES)


def extract(path):








    path = Path(path)
    try:
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER
    except ImportError:
        return Document(path=str(path), format=FORMAT, units=[],
                        cannot_see=[office.missing_package(PACKAGE, FORMAT)], encoding="utf-8",
                        newline="\n")
    try:
        deck = Presentation(str(path))
    except Exception as error:
        raise office.OfficeError(f"{path}: {PACKAGE} could not open the file, "
                                 f"{type(error).__name__}: {error}") from error
    walk = _Walk(titles={PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE},
                 furniture={PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE,
                            PP_PLACEHOLDER.FOOTER})
    for number, slide in enumerate(deck.slides, start=1):
        walk.slide(slide, number)
    unreached = walk.on_slides - walk.reached
    if unreached > 0:
        walk.lines.append(f"{unreached} text nodes on the slides were not reached by the walk")
    passed_over = _template_text(deck)
    if passed_over:
        walk.lines.append(f"{passed_over} text nodes in the slide layouts and masters were "
                          "passed over, since the walk reads the slides alone")
    return Document(path=str(path), format=FORMAT, units=walk.units,
                    cannot_see=list(CANNOT_SEE) + walk.lines, encoding="utf-8", newline="\n",
                    protected=walk.protected)



FIELD = "the paragraph holds a field (a:fld), whose text a rewrite would move or lose"
EMPTIED = "the rewrite empties the paragraph"
CHANGED = ("the paragraph's runs, breaks and fields differ from the ones the extract read, so "
           "it changed between extract and write")
ABSENT = "the paragraph the locator names is absent from the deck"


PART_ABSENT = ("no part of the deck holds the paragraph of {surface}, so the in-place write "
               "could not name the entry to replace")

_SURROGATES = (chr(0xD800), chr(0xDFFF))
_NONCHARACTERS = (chr(0xFFFE), chr(0xFFFF))


def _paragraph_at(slide, shapes, unit):




    locator = unit.locator
    if locator.get("notes"):
        placeholder = slide.notes_slide.notes_placeholder if slide.has_notes_slide else None
        body = None if placeholder is None else placeholder.element.find(_P + "txBody")
    else:
        shape = shapes.get(tuple(locator["shape"]))
        body = None if shape is None else shape.element.find(_P + "txBody")
        if shape is not None and "cell" in locator:
            row, column = locator["cell"]
            grid = shape.element.find(".//" + _A + "tbl")
            lines = [] if grid is None else grid.findall(_A + "tr")
            cells = lines[row].findall(_A + "tc") if row < len(lines) else []
            body = cells[column].find(_A + "txBody") if column < len(cells) else None
    paragraphs = [] if body is None else body.findall(_A + "p")
    index = locator["paragraph"]
    return paragraphs[index] if index < len(paragraphs) else None


def _lines(element):

    lines = [[]]
    for child in element:
        if child.tag == _A + "r":
            lines[-1].append(child)
        elif child.tag == _A + "br":
            lines.append([])
    return lines


def _rewrite(element, unit, text):









    read, runs, breaks, fields = _paragraph(element)
    frame = unit.frame
    if (read, runs, breaks, len(fields)) != (unit.text, frame["runs"], frame["breaks"],
                                             frame["fields"]):
        return CHANGED
    parts = text.replace("\r\n", "\n").split("\n")
    if len(parts) != breaks + 1:
        return (f"the rewrite holds {len(parts) - 1} line breaks where the paragraph holds "
                f"{breaks} a:br")
    pieces = []
    for number, (line, part) in enumerate(zip(_lines(element), parts), start=1):
        if not line and part:
            return (f"the rewrite puts text on line {number} of the paragraph, which holds "
                    "no run")
        if line:
            pieces.extend(zip(line, office.distribute([len(run.text) for run in line], part)))
    for run, piece in pieces:
        run.text = piece
    return None


def _refused_before_reading(unit, text):







    if not unit.frame.get("writable"):
        return office.READ_ONLY
    if unit.frame.get("fields"):
        return FIELD
    if not text.strip():
        return EMPTIED
    control = next((char for char in text.replace("\r\n", "\n")
                    if char < " " and char not in "\t\n"), None)
    if control is not None:
        return (f"the rewrite holds the control character U+{ord(control):04X}, which "
                "python-pptx would write into the slide as an escape")
    first, last = _SURROGATES
    held = next((char for char in text if first <= char <= last or char in _NONCHARACTERS),
                None)
    if held is not None:
        return (f"the rewrite holds the character U+{ord(held):04X}, which cannot be "
                "written into an a:t")
    return None


def _apply(deck, document, texts):









    applied, refused, lines, written = [], [], [], []
    slides, shapes = list(deck.slides), {}
    for unit in document.units:
        if unit.id not in texts:
            continue
        reason = _refused_before_reading(unit, texts[unit.id])
        if reason is None:
            number, element = unit.locator["slide"], None
            if 1 <= number <= len(slides):
                slide = slides[number - 1]
                if number not in shapes:
                    shapes[number] = dict(_shapes(slide.shapes))
                element = _paragraph_at(slide, shapes[number], unit)
            reason = ABSENT if element is None else _rewrite(element, unit, texts[unit.id])
        if reason is None:
            applied.append(unit.id)
            written.append(element)
            if element.find(f"{_A}r/{_A}rPr/{_A}hlinkClick") is not None:
                lines.append(f"the rewrite of {unit.surface} was spread over runs that "
                             "include a hyperlink, so the words under the link may have "
                             "moved with the rewrite")
            if any(office.splits_a_word([run.text for run in line])
                   for line in _lines(element)):
                lines.append(office.SPLIT_WORD.format(surface=unit.surface))
        else:
            refused.append((unit.id, reason))
    return applied, refused, lines, written


def write_copy(path, document, texts, file_digest):

















    path = Path(path)
    target = office.copy_path(path)
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        from pptx import Presentation
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    try:
        deck = Presentation(io.BytesIO(data))
        applied, refused, lines, _ = _apply(deck, document, texts)
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if texts and not applied:
        return None, [], refused
    try:
        office.save_copy(target, deck.save)
    except Exception as error:
        return None, [], [(None, f"the copy could not be saved at {target}: "
                                 f"{type(error).__name__}: {error}")]
    document.cannot_see.extend(lines)
    return target, applied, refused


def _part_names(deck):





    return {id(part): str(part.partname).lstrip("/")
            for part in deck.part.package.iter_parts()}


def _owned_part(deck, element):



    root = element.getroottree().getroot()
    return next((part for part in deck.part.package.iter_parts()
                 if getattr(part, "_element", None) is root), None)


def write_in_place(path, document, texts, file_digest):




















    path = Path(path)
    if not texts:
        return None, [], []
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        from pptx import Presentation
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    try:
        deck = Presentation(io.BytesIO(data))
        names = _part_names(deck)
        applied, refused, lines, written = _apply(deck, document, texts)
        replacements = {}
        surfaces = {unit.id: unit.surface for unit in document.units}
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            order = package.namelist()
            for unit_id, element in zip(applied, written):
                part = _owned_part(deck, element)
                if part is None:
                    return None, [], [(None, PART_ABSENT.format(surface=surfaces[unit_id]))]
                name = names[id(part)]
                if name not in replacements:
                    replacements[name] = office.with_leading_bytes(package.read(name),
                                                                   part.blob)
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if not applied:
        return None, [], refused
    kept = {unit_id: texts[unit_id] for unit_id in applied}

    def census(temp):
        return office.reads_as(document, extract(temp), kept)

    reason = office.write_in_place(path, data, replacements, census)
    if reason is not None:
        return None, [], [(None, reason)]
    replaced = [name for name in order if name in replacements]
    document.cannot_see.extend(lines)
    document.cannot_see.append(office.IN_PLACE_PARTS.format(
        path=path, count=len(replaced), total=len(order), names=", ".join(replaced)))
    return path, applied, refused
