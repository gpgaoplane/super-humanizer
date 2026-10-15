






















































































import io
import zipfile
from collections import Counter
from pathlib import Path

from scripts import office
from scripts.records import Document, Unit

FORMAT = "docx"
PACKAGE = "python-docx"


FLOOR = "1.1"


CANNOT_SEE = (
    "a text box or a shape's text",
    "a content control's text",
    "text inside a tracked deletion",
    "a comment part",
    "SmartArt",
    "a field's instruction text, which is read through its last result alone",
    "a hyperlink's target",
    "an image and any text inside it",
    "footnote and endnote text, which is read and counted and never rewritten in this "
    "version",
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_P, _R, _T = _W + "p", _W + "r", _W + "t"






_SEPARATORS = {_W + "tab": ("tab", chr(9)), _W + "ptab": ("ptab", chr(9)),
               _W + "br": ("break", chr(10)), _W + "cr": ("cr", chr(10)),
               _W + "noBreakHyphen": ("hyphen", "-")}









_BLOCKING = {**{_W + name: "w:" + name for name in (
    "drawing", "pict", "sym", "ruby", "footnoteReference", "endnoteReference",
    "commentReference", "fldChar", "instrText", "fldSimple", "pgNum", "dayShort", "dayLong",
    "monthShort", "monthLong", "yearShort", "yearLong", "noBreakHyphen", "softHyphen", "ins",
    "del", "moveFrom", "moveTo", "sdt")}, _MC + "AlternateContent": "mc:AlternateContent"}



_DELETED = frozenset((_W + "del", _W + "moveFrom"))



_PASSED_OVER = frozenset((_W + "pPr", _W + "sdt", _MC + "AlternateContent"))



_HEADERS = (("header", "header", "default", "header"),
            ("first_page_header", "header", "first", "first page header"),
            ("even_page_header", "header", "even", "even page header"),
            ("footer", "footer", "default", "footer"),
            ("first_page_footer", "footer", "first", "first page footer"),
            ("even_page_footer", "footer", "even", "even page footer"))



_NOTES = (("FOOTNOTES", "footnotes", "footnote"), ("ENDNOTES", "endnotes", "endnote"))



_C141 = frozenset(("quote", "intensequote", "blocktext", "bibliography"))


_KIND_WORDS = {"paragraph": "paragraph", "heading": "heading", "list_item": "list item",
               "caption": "caption"}


def _runs(element, tracked):




    for child in element:
        if child.tag == _R:
            yield from child
        elif child.tag in _DELETED:
            tracked.append(child)
        elif child.tag not in _PASSED_OVER:
            yield from _runs(child, tracked)


def _nodes(element, tracked):

    for node in _runs(element, tracked):
        if node.tag == _T or node.tag in _SEPARATORS:
            yield node


def nodes(paragraph):



    return list(_nodes(paragraph, []))


def _read(element):




    pieces, texts, separators, cuts, tracked = [], [], [], [], []
    reached = 0
    for node in _nodes(element, tracked):
        if node.tag == _T:
            text = node.text or ""
            pieces.append(text)
            texts.append(len(text))
            reached += bool(text.strip())
        else:
            name, rendered = _SEPARATORS[node.tag]
            pieces.append(rendered)
            separators.append(name)
            cuts.append(len(texts))
    blocked = list(dict.fromkeys(_BLOCKING[node.tag] for child in element
                                 if child.tag != _W + "pPr" for node in child.iter()
                                 if node.tag in _BLOCKING))
    return "".join(pieces), texts, separators, cuts, blocked, len(tracked), reached


def _styles(part, styles_type):


    names, default = {}, None
    for rel in part.rels.values():
        if rel.reltype != styles_type or rel.is_external:
            continue
        for style in rel.target_part.element.iter(_W + "style"):
            if style.get(_W + "type", "paragraph") != "paragraph":
                continue
            sid = style.get(_W + "styleId")
            name = style.find(_W + "name")
            names[sid] = sid if name is None else name.get(_W + "val", sid)
            if style.get(_W + "default") in ("1", "true", "on"):
                default = sid
    return names, default


def _style(element, styles):



    names, default = styles
    node = element.find(f"{_W}pPr/{_W}pStyle")
    sid = default if node is None else node.get(_W + "val")
    return (names.get(sid, sid) or "").replace(" ", "").lower()


def _numbered(element):


    numbering = element.find(f"{_W}pPr/{_W}numPr")
    if numbering is None:
        return False
    number = numbering.find(_W + "numId")
    return number is None or number.get(_W + "val") != "0"


def _kind(element, style):

    if style == "title" or style.startswith("heading"):
        return "heading"
    if _numbered(element) or style.startswith("list"):
        return "list_item"
    if style == "caption":
        return "caption"
    return "paragraph"


class _Walk:


    def __init__(self, styles, table_class, cell_class):
        self.units, self.protected, self.lines = [], [], []
        self.styles, self.table_class, self.cell_class = styles, table_class, cell_class
        self.tracked = self.reached = self.in_parts = 0
        self.part({}, {})

    def part(self, positions, base, prefix="", kind=None, writable=True):



        self.positions, self.base, self.prefix = positions, base, prefix
        self.kind, self.writable = kind, writable
        self.counts, self.tables = Counter(), 0

    def _next(self, key):

        self.counts[key] += 1
        return self.counts[key]

    def blocks(self, container, cell=None):


        for item in container.iter_inner_content():
            if isinstance(item, self.table_class):
                self._table(item)
            else:
                self.paragraph(item._p, cell)

    def _table(self, table):
        self.tables += 1
        number = self.tables
        for row, line in enumerate(table._tbl.tr_lst):
            for column, cell in enumerate(line.tc_lst):
                self.blocks(self.cell_class(cell, table), (number, row, column))

    def paragraph(self, element, cell=None):

        text, texts, separators, cuts, blocked, tracked, reached = _read(element)
        self.tracked += tracked
        if not text.strip():
            return
        self.reached += reached
        style = _style(element, self.styles)
        locator = dict(self.base)
        if cell is None:
            kind = self.kind or _kind(element, style)
            word = "paragraph" if self.kind else _KIND_WORDS[kind]
            surface = f"{self.prefix}{word} {self._next(word)}"
        else:
            number, row, column = cell
            kind = self.kind or "table_cell"
            locator.update(table=number - 1, row=row, col=column)
            surface = (f"{self.prefix}table {number} cell r{row + 1}c{column + 1} "
                       f"paragraph {self._next(cell)}")
        locator["paragraph"] = self.positions[element]
        quoted = style in _C141
        unit = Unit(
            id=f"u{len(self.units) + 1:04d}", kind=kind, surface=surface, text=text,
            locator=locator, weighs=not quoted,
            frame={"kind": "docx_paragraph", "texts": texts, "separators": separators,
                   "cuts": cuts, "blocked": blocked, "writable": self.writable})
        self.units.append(unit)
        if quoted:
            self.protected.append((unit.id, 0, len(text), "C141"))

    def notes(self, root, part, word):



        positions = _positions(root)
        self.in_parts += _text_nodes(root)
        for note in root.iter(_W + word):
            try:
                number = int(note.get(_W + "id"))
            except (TypeError, ValueError):
                self.lines.append(f"a {word} whose w:id is missing or not a whole number was "
                                  "skipped, so its text was not read")
                continue
            self.part(positions, {"part": part, "note": number}, prefix=f"{word} {number} ",
                      kind="notes", writable=False)
            for element in _note_paragraphs(note):
                self.paragraph(element)

    def alt_texts(self, root, base, prefix):


        for index, drawing in enumerate(root.iter(_W + "drawing")):
            properties = next(drawing.iter(_WP + "docPr"), None)
            text = "" if properties is None else properties.get("descr") or ""
            if not text.strip():
                continue
            self.units.append(Unit(
                id=f"u{len(self.units) + 1:04d}", kind="label",
                surface=f"{prefix}picture {index + 1} alt text", text=text,
                locator={**base, "drawing": index},
                frame={"kind": "alt_text", "writable": False}))


def _note_paragraphs(element):

    for child in element:
        if child.tag == _P:
            yield child
        elif child.tag in (_W + "tbl", _W + "tr", _W + "tc"):
            yield from _note_paragraphs(child)


def _positions(root):

    return {element: index for index, element in enumerate(root.iter(_P))}


def _text_nodes(element):





    count = 0
    for child in element:
        if child.tag == _T:
            count += bool((child.text or "").strip())
        elif child.tag != _MC + "Fallback" and child.tag not in _DELETED:
            count += _text_nodes(child)
    return count


def extract(path):











    path = Path(path)
    try:
        import docx
        from docx.document import Document as WordDocument
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml.parser import parse_xml
        from docx.table import Table, _Cell
    except ImportError:
        return Document(path=str(path), format=FORMAT, units=[],
                        cannot_see=[office.missing_package(PACKAGE, FORMAT)], encoding="utf-8",
                        newline="\n")
    if not hasattr(WordDocument, "iter_inner_content"):
        return Document(path=str(path), format=FORMAT, units=[],
                        cannot_see=[office.below_floor(PACKAGE, FLOOR, FORMAT)],
                        encoding="utf-8", newline="\n")
    try:
        document = docx.Document(str(path))
    except Exception as error:
        raise office.OfficeError(f"{path}: {PACKAGE} could not open the file, "
                                 f"{type(error).__name__}: {error}") from error
    if document.element.body is None:
        raise office.OfficeError(f"{path}: {PACKAGE} opened the file, and the document part "
                                 "holds no w:body")
    walk = _Walk(_styles(document.part, RELATIONSHIP_TYPE.STYLES), Table, _Cell)
    walk.part(_positions(document.element), {"part": "body"})
    walk.blocks(document)
    walk.in_parts += _text_nodes(document.element)
    drawn, walked = [(document.element, {"part": "body"}, "")], set()
    for number, section in enumerate(document.sections):
        for attribute, part, which, words in _HEADERS:
            header = getattr(section, attribute)
            if header.is_linked_to_previous or header.part.partname in walked:
                continue
            walked.add(header.part.partname)
            root, base = header.part.element, {"part": part, "section": number, "which": which}
            prefix = f"section {number + 1} {words} "
            walk.part(_positions(root), base, prefix=prefix, kind="label")
            walk.blocks(header)
            walk.in_parts += _text_nodes(root)
            drawn.append((root, base, prefix))
    for name, part, word in _NOTES:
        for rel in document.part.rels.values():
            if rel.reltype != getattr(RELATIONSHIP_TYPE, name) or rel.is_external:
                continue
            try:
                root = parse_xml(rel.target_part.blob)
            except SyntaxError:
                walk.lines.append(f"the {part} part could not be parsed, so its notes were "
                                  "not read")
                continue
            walk.notes(root, part, word)
    for root, base, prefix in drawn:
        walk.alt_texts(root, base, prefix)
    if walk.tracked:
        walk.lines.append(f"{walk.tracked} tracked changes were passed over, so the text "
                          "inside them was not read")
    unreached = walk.in_parts - walk.reached
    if unreached > 0:
        walk.lines.append(f"{unreached} text nodes in the parts the walk reads were not "
                          "reached by it")
    return Document(path=str(path), format=FORMAT, units=walk.units,
                    cannot_see=list(CANNOT_SEE) + walk.lines, encoding="utf-8", newline="\n",
                    protected=walk.protected)



EMPTIED = "the rewrite empties the paragraph"
ABSENT = "the paragraph the locator names is absent from the document"
CHANGED = ("the paragraph's text differs from the text the extract read, so it changed "
           "between extract and write")
ORDER = ("the rewrite's tabs and line breaks come in another order than the paragraph's "
         "separators")



PART_ABSENT = ("no part of the document holds the paragraph of {surface}, so the in-place "
               "write could not name the entry to replace")




LINKED = ("the rewrite of {surface} was spread over runs that include a hyperlink, so the "
          "words under the link may have moved with the rewrite")
UNSHOWN = ("the rewrite of {surface} was written into a header or footer that Word shows only "
           "when {setting} is on, and this document leaves it off")
_SHOWN_WHEN = {"first": "its section's different first page (w:titlePg)",
               "even": "the settings' different odd and even pages (w:evenAndOddHeaders)"}


_RENDERED = dict(_SEPARATORS.values())



_HEADER_ATTRIBUTES = {(part, which): attribute for attribute, part, which, _ in _HEADERS}

_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_TAB, _NEWLINE, _RETURN = chr(9), chr(10), chr(13)
_SURROGATES = (chr(0xD800), chr(0xDFFF))
_NONCHARACTERS = (chr(0xFFFE), chr(0xFFFF))


def _blocked(names):

    return (f"the paragraph holds {', '.join(names)}, which a rewrite of its w:t nodes would "
            "break or leave beside the new text")


def _shape(texts, separators):

    named = "the separators " + ", ".join(separators) if separators else "no separator"
    return f"{len(texts)} w:t and {named}"


def _unwritable(char):



    first, last = _SURROGATES
    return ((char < " " and char not in (_TAB, _NEWLINE)) or first <= char <= last
            or char in _NONCHARACTERS)


def _refused_before_reading(unit, text):





    frame = unit.frame
    if not frame.get("writable") or frame.get("kind") != "docx_paragraph":
        return office.READ_ONLY
    if frame["blocked"]:
        return _blocked(frame["blocked"])
    if not text.strip():
        return EMPTIED
    held = next((char for char in text.replace(_RETURN + _NEWLINE, _NEWLINE)
                 if _unwritable(char)), None)
    if held is not None:
        kind = "control character" if held < " " else "character"
        return (f"the rewrite holds the {kind} U+{ord(held):04X}, which cannot be written "
                "into a w:t")
    return None


def _part_paragraphs(opened, locator):




    if locator.get("part") == "body":
        return list(opened.element.iter(_P))
    attribute = _HEADER_ATTRIBUTES.get((locator.get("part"), locator.get("which")))
    sections, number = opened.sections, locator.get("section")
    if attribute is None or not isinstance(number, int) or not 0 <= number < len(sections):
        return []
    header = getattr(sections[number], attribute)
    return [] if header.is_linked_to_previous else list(header.part.element.iter(_P))


def _paragraph_at(opened, parts, locator):



    key = (locator.get("part"), locator.get("section"), locator.get("which"))
    if key not in parts:
        parts[key] = _part_paragraphs(opened, locator)
    paragraphs, index = parts[key], locator.get("paragraph")
    return paragraphs[index] if isinstance(index, int) and 0 <= index < len(paragraphs) else None


def _on(element, tag):


    node = element.find(tag)
    return node is not None and node.get(_W + "val", "true") in ("1", "true", "on")


def _shown(opened, locator, settings_type):





    which = locator.get("which")
    if which == "first":
        return opened.sections[locator["section"]].different_first_page_header_footer
    if which == "even":
        return any(_on(rel.target_part.element, _W + "evenAndOddHeaders")
                   for rel in opened.part.rels.values()
                   if rel.reltype == settings_type and not rel.is_external)
    return True


def _owed(opened, unit, element, settings_type):

    lines = []
    if any(True for link in element.iter(_W + "hyperlink") for _ in link.iter(_T)):
        lines.append(LINKED.format(surface=unit.surface))
    if any(office.splits_a_word([node.text or "" for node in group])
           for group in _segments(element)):
        lines.append(office.SPLIT_WORD.format(surface=unit.surface))
    if unit.locator.get("part") != "body" and not _shown(opened, unit.locator, settings_type):
        lines.append(UNSHOWN.format(surface=unit.surface,
                                    setting=_SHOWN_WHEN[unit.locator["which"]]))
    return lines


def _segments(element):


    segments = [[]]
    for node in nodes(element):
        if node.tag == _T:
            segments[-1].append(node)
        else:
            segments.append([])
    return segments


def _rewrite(element, unit, text):












    frame = unit.frame
    read, texts, separators, cuts, blocked, _, _ = _read(element)
    if blocked:
        return _blocked(blocked)
    if (len(texts), separators, cuts) != (len(frame["texts"]), frame["separators"],
                                          frame["cuts"]):
        return (f"the paragraph holds {_shape(texts, separators)} where the extract read "
                f"{_shape(frame['texts'], frame['separators'])}, so it changed between "
                "extract and write")
    if (read, texts) != (unit.text, frame["texts"]):
        return CHANGED
    text = text.replace(_RETURN + _NEWLINE, _NEWLINE)
    marks = [char for char in text if char in (_TAB, _NEWLINE)]
    rendered = [_RENDERED[name] for name in separators]
    if marks != rendered:
        if sorted(marks) == sorted(rendered):
            return ORDER
        return (f"the rewrite holds {marks.count(_TAB)} tabs and {marks.count(_NEWLINE)} "
                f"line breaks where the paragraph's separators render "
                f"{rendered.count(_TAB)} tabs and {rendered.count(_NEWLINE)} line breaks")
    pieces = []
    parts = text.replace(_TAB, _NEWLINE).split(_NEWLINE)
    for number, (group, part) in enumerate(zip(_segments(element), parts), start=1):
        if not group and part:
            return (f"the rewrite puts text in segment {number} of the paragraph, which "
                    "holds no w:t")
        if group:
            lengths = [len(node.text or "") for node in group]
            pieces.extend(zip(group, office.distribute(lengths, part)))
    for node, piece in pieces:
        node.text = piece
        node.set(_XML_SPACE, "preserve")
    return None


def _apply(opened, document, texts, settings_type):








    applied, refused, lines, written, parts = [], [], [], [], {}
    for unit in document.units:
        if unit.id not in texts:
            continue
        reason = _refused_before_reading(unit, texts[unit.id])
        if reason is None:
            element = _paragraph_at(opened, parts, unit.locator)
            reason = ABSENT if element is None else _rewrite(element, unit, texts[unit.id])
        if reason is None:
            applied.append(unit.id)
            written.append(element)
            lines.extend(_owed(opened, unit, element, settings_type))
        else:
            refused.append((unit.id, reason))
    return applied, refused, lines, written


def write_copy(path, document, texts, file_digest):


















    path = Path(path)
    target = office.copy_path(path)
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        import docx
        from docx.opc.constants import RELATIONSHIP_TYPE
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    try:
        opened = docx.Document(io.BytesIO(data))
        applied, refused, lines, _ = _apply(opened, document, texts,
                                            RELATIONSHIP_TYPE.SETTINGS)
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if texts and not applied:
        return None, [], refused
    try:
        office.save_copy(target, opened.save)
    except Exception as error:
        return None, [], [(None, f"the copy could not be saved at {target}: "
                                 f"{type(error).__name__}: {error}")]
    document.cannot_see.extend(lines)
    return target, applied, refused


def _owned_part(opened, element):




    root = element.getroottree().getroot()
    return next((part for part in opened.part.package.iter_parts()
                 if getattr(part, "element", None) is root), None)


def write_in_place(path, document, texts, file_digest):























    path = Path(path)
    if not texts:
        return None, [], []
    data, reason = office.read_unchanged(path, file_digest)
    if reason is not None:
        return None, [], [(None, reason)]
    try:
        import docx
        from docx.opc.constants import RELATIONSHIP_TYPE
    except ImportError:
        return None, [], [(None, office.missing_package(PACKAGE, FORMAT))]
    try:
        opened = docx.Document(io.BytesIO(data))
        applied, refused, lines, written = _apply(opened, document, texts,
                                                  RELATIONSHIP_TYPE.SETTINGS)
        replacements = {}
        surfaces = {unit.id: unit.surface for unit in document.units}
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            order = package.namelist()
            for unit_id, element in zip(applied, written):
                part = _owned_part(opened, element)
                if part is None:
                    return None, [], [(None, PART_ABSENT.format(surface=surfaces[unit_id]))]
                name = str(part.partname).lstrip("/")
                if name not in replacements:
                    replacements[name] = office.with_leading_bytes(package.read(name),
                                                                   part.blob)
    except Exception as error:
        return None, [], [(None, f"{PACKAGE} raised {type(error).__name__}: {error}")]
    if not applied:
        return None, [], refused
    kept = {unit_id: texts[unit_id] for unit_id in applied}

    def census(temp):
        return office.reads_as(document, extract(temp), kept)

    reason = office.write_in_place(path, data, replacements, census)
    if reason is not None:
        return None, [], [(None, reason)]
    replaced = [name for name in order if name in replacements]
    document.cannot_see.extend(lines)
    document.cannot_see.append(office.IN_PLACE_PARTS.format(
        path=path, count=len(replaced), total=len(order), names=", ".join(replaced)))
    return path, applied, refused
